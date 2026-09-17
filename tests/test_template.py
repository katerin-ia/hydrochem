"""Las plantillas de Excel se leen con el mismo lector que el resto.

La prueba que importa no es que el archivo exista, sino que HydroChem sea capaz
de leer lo que el mismo reparte: si un dia cambian los nombres que reconoce el
mapeo automatico y la plantilla se queda atras, el usuario abre el archivo que
le dimos, lo rellena y la aplicacion no le detecta las columnas. Aqui la
plantilla se genera, se lee y se comprueba el mapeo campo a campo.
"""

from __future__ import annotations

import io
import tempfile
from pathlib import Path

import openpyxl

from hydrochem.io.readers import read_table
from hydrochem.io.template import build_template, filename
from hydrochem.pipeline import AnalysisOptions, analyse

#: Lo que el nucleo necesita para trabajar. Si la plantilla deja de ofrecer uno
#: de estos, la aplicacion no puede hacer ni el Piper.
IMPRESCINDIBLES = [
    "station_code", "Ca_mgL", "Mg_mgL", "Na_mgL", "K_mgL",
    "HCO3_mgL", "SO4_mgL", "Cl_mgL",
]


def _leer(kind: str):
    carpeta = Path(tempfile.mkdtemp())
    ruta = carpeta / f"{kind}.xlsx"
    ruta.write_bytes(build_template(kind))
    return read_table(ruta)


def test_hay_dos_plantillas_con_nombres_distintos():
    assert filename("simple") != filename("campaigns")
    assert "campanas" in filename("campaigns")


def test_tipo_desconocido_se_rechaza_con_mensaje_util():
    try:
        build_template("mensual")
    except ValueError as exc:
        assert "simple" in str(exc) and "campaigns" in str(exc)
    else:
        raise AssertionError("una plantilla inventada deberia fallar")


def test_la_plantilla_simple_se_mapea_entera():
    res = _leer("simple")
    for clave in IMPRESCINDIBLES:
        assert clave in res.mapping.mapping, f"falta {clave} en el mapeo"


def test_la_plantilla_de_campanas_anade_la_fecha():
    res = _leer("campaigns")
    for clave in IMPRESCINDIBLES + ["sampled_at"]:
        assert clave in res.mapping.mapping, f"falta {clave} en el mapeo"


def test_la_fila_de_ayuda_no_se_cuela_como_dato():
    # La fila 2 explica cada columna. Si el lector la tomara por una muestra,
    # el usuario veria una fila de texto entre sus datos.
    res = _leer("simple")
    assert len(res.data) == 1, "solo deberia quedar la fila de ejemplo"


def test_el_ejemplo_de_campanas_repite_estaciones_en_varias_fechas():
    """Es el punto entero de esta plantilla: ensenar el formato largo.

    Si el ejemplo trajera una fila por estacion, no enseniaria nada que la
    plantilla simple no ensene ya.
    """
    res = _leer("campaigns")
    ds = analyse(res.data, AnalysisOptions(), mapping=res.mapping)
    assert ds.has_dates
    por_estacion = ds.data.groupby("station_code")["sampled_at"].nunique()
    assert por_estacion.min() >= 3, "cada estacion necesita varias campanas"
    assert len(por_estacion) >= 2, "hace falta mas de una estacion"


def test_el_ejemplo_no_llega_con_el_balance_roto():
    """Un ejemplo con aviso de balance ensena a ignorar los avisos."""
    for kind in ("simple", "campaigns"):
        res = _leer(kind)
        ds = analyse(res.data, AnalysisOptions(), mapping=res.mapping)
        assert (ds.data["CBE_pct"].abs() <= 5.0).all(), kind


def test_lleva_hoja_de_instrucciones():
    for kind in ("simple", "campaigns"):
        wb = openpyxl.load_workbook(io.BytesIO(build_template(kind)))
        assert "LEEME" in wb.sheetnames
        assert "Datos" in wb.sheetnames
        texto = "\n".join(
            str(c.value or "") for (c,) in wb["LEEME"].iter_rows(max_col=1)
        )
        # El aviso sobre los ceros es el que evita el error mas caro: un 0
        # escrito donde no se midio cambia el balance y las excedencias.
        assert "VACIA" in texto or "vacia" in texto
        if kind == "campaigns":
            assert "columna por campana" in texto
