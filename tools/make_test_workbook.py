"""Genera un Excel de PRUEBA para ejercitar la aplicacion entera.

Por que existe: el libro de referencia (Amargosa) no tiene fechas, ni pH, ni
coordenadas UTM, ni valores censurados. Con el no se pueden probar la mitad de
las pantallas. Este archivo los trae todos, y esta hecho para que cada funcion
de la aplicacion se encienda:

=========================  ====================================================
Que prueba                 Como
=========================  ====================================================
Mapeo de columnas          Cabeceras en castellano, con acentos y plurales
Coordenadas UTM            Este/Norte en UTM 18S (Lima Sur), no grados
Modulo temporal            3 campanas semestrales de las mismas 15 estaciones
Limites de deteccion       Celdas con ``<0,05`` y ``ND``
Datos ausentes             Algunas celdas vacias
Paneles del Durov          Trae pH y temperatura
Facies variadas            Tres zonas: bicarbonatada, mixta y clorurada
Umbrales normativos        Fluoruro y nitrato por encima del limite en algunos puntos
Avisos de validacion       Un valor fuera de rango a proposito
=========================  ====================================================

**Los datos son inventados.** Las concentraciones se construyen para que el
balance de carga cierre (es lo que hace que el archivo sea util como prueba), no
para representar ningun acuifero real. La primera hoja del libro lo advierte.

Uso:
    python tools/make_test_workbook.py
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "core"))

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from hydrochem.constants import IUPAC2021
from hydrochem.geo.crs import geographic_to_utm, utm_crs

DESTINO = ROOT / "data" / "samples" / "EJEMPLO-PRUEBA-LimaSur.xlsx"

CAMPANAS = [("2024-03-15", "C1"), ("2024-09-20", "C2"), ("2025-03-18", "C3")]

#: Tres zonas con quimica distinta, para que las facies salgan variadas.
#: Los pesos son fracciones de la suma de cationes y de aniones en meq/L.
ZONAS = {
    "Acuifero superior": {
        "n": 6,
        "meq_total": (3.2, 5.0),
        "cat": (0.52, 0.18, 0.27, 0.03),   # Ca, Mg, Na+K reparto
        "an": (0.72, 0.16, 0.12),          # HCO3, SO4, Cl
        "ph": (7.0, 7.6),
        "temp": (19.5, 22.0),
    },
    "Zona de transicion": {
        "n": 5,
        "meq_total": (6.0, 10.0),
        "cat": (0.34, 0.16, 0.46, 0.04),
        "an": (0.42, 0.24, 0.34),
        "ph": (7.2, 7.9),
        "temp": (20.0, 23.5),
    },
    "Cuna salina": {
        "n": 4,
        "meq_total": (12.0, 26.0),
        "cat": (0.18, 0.14, 0.64, 0.04),
        "an": (0.16, 0.18, 0.66),
        "ph": (7.4, 8.2),
        "temp": (21.0, 24.5),
    },
}

CABECERAS = [
    ("Fecha", ""), ("Campaña", ""), ("Estación", ""), ("Grupo", ""),
    ("Este", "m, UTM 18S"), ("Norte", "m, UTM 18S"),
    ("pH", "u. de pH"), ("Temperatura", "°C"),
    ("Conductividad", "µS/cm"), ("STD", "mg/L"),
    ("Calcio", "mg/L"), ("Magnesio", "mg/L"), ("Sodio", "mg/L"), ("Potasio", "mg/L"),
    ("Bicarbonatos", "mg/L"), ("Carbonatos", "mg/L"), ("Sulfatos", "mg/L"),
    ("Cloruros", "mg/L"), ("Fluoruros", "mg/L"), ("Nitratos", "mg/L"),
]

AVISO = (
    "DATOS INVENTADOS — ARCHIVO DE PRUEBA",
    "",
    "Este libro se genero para probar la aplicacion HydroChem. Las",
    "concentraciones estan construidas para que el balance de carga cierre",
    "y para que cada funcion de la herramienta se pueda ejercitar.",
    "",
    "NO representa ningun acuifero real y no sirve para sacar ninguna",
    "conclusion hidroquimica ni para ningun informe.",
    "",
    "Que trae, y para que:",
    "  · Cabeceras en castellano, para probar el mapeo automatico de columnas.",
    "  · Coordenadas en UTM 18S (no en grados), para probar la reproyeccion.",
    "  · Tres campanas de las mismas 15 estaciones, para el modulo temporal.",
    "  · Valores censurados (<0,05) y celdas vacias.",
    "  · pH y temperatura, para que los paneles del Durov se dibujen.",
    "  · Un valor fuera de rango a proposito, para ver el aviso de validacion.",
    "",
    "Se regenera con:  python tools/make_test_workbook.py",
)


def _reparto(rng, pesos, total):
    """Reparte un total entre varios componentes con algo de variacion."""
    ruido = rng.normal(1.0, 0.10, len(pesos))
    crudo = np.array(pesos) * np.clip(ruido, 0.55, 1.6)
    return crudo / crudo.sum() * total


def main() -> int:
    rng = np.random.default_rng(20260912)
    crs = utm_crs(18, "S")

    # --- estaciones: coordenadas alrededor de Lima Sur, en UTM 18S ----------
    estaciones = []
    idx = 1
    for zona, cfg in ZONAS.items():
        for _ in range(cfg["n"]):
            # Villa El Salvador y alrededores; la cuna salina, mas cerca del mar
            lon = -76.96 + rng.normal(0, 0.035)
            lat = -12.21 + rng.normal(0, 0.035)
            if zona == "Cuna salina":
                lon -= 0.045  # hacia la costa
            este, norte = geographic_to_utm(lon, lat, crs)
            estaciones.append({
                "codigo": f"VES-{idx:02d}",
                "zona": zona,
                "este": round(float(este), 1),
                "norte": round(float(norte), 1),
                "base_meq": rng.uniform(*cfg["meq_total"]),
            })
            idx += 1

    filas = []
    for k, (fecha, campana) in enumerate(CAMPANAS):
        for est in estaciones:
            cfg = ZONAS[est["zona"]]
            # Deriva: la cuna salina avanza campana a campana
            factor = 1.0 + (0.16 * k if est["zona"] == "Cuna salina" else 0.03 * k)
            total_cat = est["base_meq"] * factor * rng.normal(1.0, 0.03)

            ca, mg, nak, k_frac = _reparto(rng, cfg["cat"], total_cat)
            # El potasio se separa del sodio: en aguas naturales es minoritario
            na = nak * 0.92
            pot = nak * 0.08 + k_frac

            # Fluoruro y nitrato se deciden ANTES de repartir los aniones
            # mayoritarios, y su aporte se descuenta del presupuesto. Si no, un
            # nitrato de 90 mg/L (1,5 meq/L) desbalancea la muestra y la
            # aplicacion la rechaza con razon.
            if est["zona"] == "Acuifero superior":
                f_txt, f_meq = "<0,05", 0.0          # censurado: no aporta carga
            elif est["zona"] == "Zona de transicion":
                f_val = round(rng.uniform(1.2, 2.6), 2)
                f_txt, f_meq = f_val, f_val / IUPAC2021["F"]
            else:
                f_val = round(rng.uniform(0.4, 1.1), 2)
                f_txt, f_meq = f_val, f_val / IUPAC2021["F"]

            if est["codigo"] in ("VES-03", "VES-07", "VES-08"):
                n_val = round(rng.uniform(52, 95), 1)
                n_txt, n_meq = n_val, n_val / IUPAC2021["NO3"]
            elif est["codigo"] in ("VES-01", "VES-12"):
                n_txt, n_meq = "ND", 0.0             # no detectado: no aporta
            else:
                n_val = round(rng.uniform(2, 28), 1)
                n_txt, n_meq = n_val, n_val / IUPAC2021["NO3"]

            # Los aniones mayoritarios cierran el balance con un desajuste
            # pequeno y aleatorio: asi el CBE sale realista, no exactamente 0.
            total_an = (ca + mg + na + pot) * rng.uniform(0.978, 1.022)
            resto = max(total_an - f_meq - n_meq, total_an * 0.35)
            hco3, so4, cl = _reparto(rng, cfg["an"], resto)

            fila = {
                "Fecha": fecha,
                "Campaña": campana,
                "Estación": est["codigo"],
                "Grupo": est["zona"],
                "Este": est["este"],
                "Norte": est["norte"],
                "pH": round(rng.uniform(*cfg["ph"]), 2),
                "Temperatura": round(rng.uniform(*cfg["temp"]), 1),
                "Calcio": round(ca * IUPAC2021["Ca"], 2),
                "Magnesio": round(mg * IUPAC2021["Mg"], 2),
                "Sodio": round(na * IUPAC2021["Na"], 2),
                "Potasio": round(pot * IUPAC2021["K"], 2),
                "Bicarbonatos": round(hco3 * IUPAC2021["HCO3"], 1),
                "Carbonatos": 0.0,
                "Sulfatos": round(so4 * IUPAC2021["SO4"], 2),
                "Cloruros": round(cl * IUPAC2021["Cl"], 2),
            }
            suma = sum(fila[c] for c in ("Calcio", "Magnesio", "Sodio", "Potasio",
                                         "Bicarbonatos", "Sulfatos", "Cloruros"))
            fila["STD"] = round(suma, 0)
            fila["Conductividad"] = round(suma * 1.55, 0)

            fila["Fluoruros"] = f_txt
            fila["Nitratos"] = n_txt
            filas.append(fila)

    # --- huecos y rarezas a proposito --------------------------------------
    filas[4]["Magnesio"] = None            # celda vacia
    filas[19]["pH"] = None                 # falta el pH de una muestra
    filas[31]["Carbonatos"] = 12.4         # algo de carbonato en una muestra
    filas[7]["Sulfatos"] = 7400.0          # fuera de rango, para ver el aviso

    # --- escritura ---------------------------------------------------------
    wb = Workbook()

    nota = wb.active
    nota.title = "LEEME"
    for i, linea in enumerate(AVISO, start=1):
        c = nota.cell(i, 1, linea)
        if i == 1:
            c.font = Font(bold=True, size=13, color="9C0006")
    nota.column_dimensions["A"].width = 78

    ws = wb.create_sheet("Datos")
    relleno = PatternFill("solid", fgColor="0E6B75")
    borde = Border(*[Side(style="thin", color="D5DFE2")] * 4)

    for col, (nombre, unidad) in enumerate(CABECERAS, start=1):
        etiqueta = f"{nombre} ({unidad})" if unidad else nombre
        celda = ws.cell(1, col, etiqueta)
        celda.fill = relleno
        celda.font = Font(bold=True, color="FFFFFF", size=10)
        celda.alignment = Alignment(horizontal="center", wrap_text=True)
        celda.border = borde
        ws.column_dimensions[get_column_letter(col)].width = max(12, len(etiqueta) * 0.75)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "C2"

    for r, fila in enumerate(filas, start=2):
        for col, (nombre, _) in enumerate(CABECERAS, start=1):
            ws.cell(r, col, fila.get(nombre)).border = borde

    DESTINO.parent.mkdir(parents=True, exist_ok=True)
    wb.save(DESTINO)

    print(f"escrito: {DESTINO}")
    print(f"  {len(filas)} filas = {len(estaciones)} estaciones x {len(CAMPANAS)} campanas")
    print(f"  zonas: {', '.join(ZONAS)}")
    print(f"  coordenadas: UTM 18S (EPSG:{crs.epsg})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
